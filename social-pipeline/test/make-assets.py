"""Creates synthetic test files in ./assets: photos, style references, a logo,
and briefing documents (PDF, DOCX, PPTX) with embedded images.
Needs Pillow, reportlab, python-docx and python-pptx."""
import math, os, random, urllib.request
from PIL import Image, ImageDraw, ImageFilter, ImageFont
os.makedirs("assets", exist_ok=True)
random.seed(4)
def font(name, size):
    for p in (f"/usr/share/fonts/truetype/dejavu/{name}.ttf",):
        if os.path.exists(p): return ImageFont.truetype(p, size)
    return ImageFont.load_default()
W, H = 3000, 2000
im = Image.new("RGB", (W, H)); d = ImageDraw.Draw(im)
for y in range(H):
    t = y / H; d.line([(0, y), (W, y)], fill=(int(170 + 60 * t), int(200 + 30 * t), int(230 - 40 * t)))
for i, (base, col) in enumerate([(900, (92, 128, 84)), (1150, (64, 100, 60)), (1400, (44, 74, 42))]):
    pts = [(0, H)] + [(x, base + int(120 * math.sin(x / (380 + i * 90) + i) + 60 * math.sin(x / 170 + i * 2))) for x in range(0, W + 50, 50)] + [(W, H)]
    d.polygon(pts, fill=col)
d.line([(1400, 2000), (1900, 1100), (2600, 700)], fill=(70, 48, 30), width=40)
for _ in range(60):
    cx, cy, r = random.randint(1500, 2700), random.randint(700, 1500), random.randint(28, 48)
    d.ellipse([cx - r, cy - r, cx + r, cy + r], fill=(random.randint(170, 210), random.randint(20, 50), random.randint(20, 40)))
im.filter(ImageFilter.GaussianBlur(1.5)).save("assets/farm-gurue.jpg", quality=90)
im = Image.new("RGB", (2400, 3000), (40, 28, 22)); d = ImageDraw.Draw(im)
d.rectangle([0, 1900, 2400, 3000], fill=(120, 80, 50))
d.ellipse([800, 1500, 1600, 1950], fill=(245, 240, 230)); d.ellipse([900, 1550, 1500, 1850], fill=(90, 50, 25))
im.filter(ImageFilter.GaussianBlur(2)).save("assets/barista-cup.jpg", quality=90)
R = Image.new("RGB", (1080, 1350), (232, 90, 40)); d = ImageDraw.Draw(R)
d.polygon([(0, 800), (1080, 500), (1080, 1350), (0, 1350)], fill=(20, 20, 20)); d.ellipse([600, 120, 1000, 520], fill=(250, 235, 210))
d.text((60, 880), "BOLD\nTYPE", font=font("DejaVuSans-Bold", 150), fill=(250, 235, 210)); R.save("assets/ref-poster.jpg", quality=88)
R = Image.new("RGB", (1080, 1350), (243, 230, 211)); d = ImageDraw.Draw(R)
d.rounded_rectangle([180, 120, 900, 860], radius=360, fill=(120, 90, 70)); d.rectangle([180, 490, 900, 860], fill=(120, 90, 70))
d.text((180, 920), "Serif\nheadline", font=font("DejaVuSerif-Bold", 110), fill=(43, 26, 18)); R.save("assets/ref-arch.png")
L = Image.new("RGBA", (900, 300), (0, 0, 0, 0)); d = ImageDraw.Draw(L)
d.ellipse([10, 40, 230, 260], fill=(43, 26, 18, 255)); d.text((260, 70), "KULUKA", font=font("DejaVuSans-Bold", 130), fill=(43, 26, 18, 255)); L.save("assets/kuluka-logo.png")
p = Image.new("RGB", (1600, 1200), (230, 220, 205)); d = ImageDraw.Draw(p)
d.rounded_rectangle([550, 250, 1050, 1050], radius=40, fill=(43, 26, 18)); d.rectangle([600, 500, 1000, 800], fill=(200, 116, 46)); p.save("assets/product-bag.jpg", quality=90)
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfgen import canvas
c = canvas.Canvas("assets/brief-kuluka.pdf", pagesize=A4); Wp, Hp = A4
c.setFont("Helvetica-Bold", 20); c.drawString(2 * cm, Hp - 3 * cm, "Briefing: Kuluka Coffee")
c.setFont("Helvetica", 11); y = Hp - 4.2 * cm
for line in ["Objectivo: dar a conhecer o primeiro café de origem única de Gurué, Zambézia.", "Público: jovens profissionais em Maputo, 25-40 anos.",
             "Mensagens: cultivado em Gurué a 1.200 m; torrado em Maputo todas as terças-feiras.", "Actividade: prova de café aos sábados às 10h00. Entrada livre, 12 lugares."]:
    c.drawString(2 * cm, y, line); y -= 0.7 * cm
c.drawImage("assets/product-bag.jpg", 2 * cm, 4 * cm, width=12 * cm, height=9 * cm); c.showPage()
c.drawImage("assets/ref-poster.jpg", 2 * cm, 8 * cm, width=8 * cm, height=10 * cm); c.save()
import docx
from docx.shared import Cm
doc = docx.Document(); doc.add_heading("Notas adicionais do cliente", 1); doc.add_paragraph("O tom deve ser caloroso e orgulhoso da origem.")
doc.add_picture("assets/product-bag.jpg", width=Cm(10)); doc.save("assets/notas-cliente.docx")
from pptx import Presentation
from pptx.util import Inches
pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "Campanha Outubro"; s.placeholders[1].text = "Três publicações."
s2 = pr.slides.add_slide(pr.slide_layouts[5]); s2.shapes.title.text = "Fotografias aprovadas"; s2.shapes.add_picture("assets/barista-cup.jpg", Inches(1), Inches(1.5), width=Inches(4))
pr.save("assets/campanha.pptx")
try:
    css = urllib.request.urlopen(urllib.request.Request("https://fonts.googleapis.com/css2?family=Lobster", headers={"User-Agent": "Mozilla/4.0"})).read().decode()
    url = css.split("url(")[1].split(")")[0]; open("assets/Lobster-Regular.ttf", "wb").write(urllib.request.urlopen(url).read())
except Exception as e:
    print("Could not download the test font:", e)
print("assets ready")
